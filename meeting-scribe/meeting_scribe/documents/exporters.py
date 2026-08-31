"""Render a :class:`DocSpec` to DOCX, ODT or PDF.

Each exporter is self-contained and depends only on its own library
(python-docx / odfpy / reportlab).
"""
from __future__ import annotations

from pathlib import Path
from typing import List

from .builders import Block, DocSpec
from ..logging_setup import get_logger

log = get_logger("export")


def export(spec: DocSpec, path: Path, fmt: str) -> Path:
    fmt = fmt.lower()
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "docx":
        _docx(spec, path)
    elif fmt == "odt":
        _odt(spec, path)
    elif fmt == "pdf":
        _pdf(spec, path)
    else:
        raise ValueError(f"Unsupported format: {fmt}")
    log.info("Wrote %s", path)
    return path


# ======================================================================= DOCX
def _docx(spec: DocSpec, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    doc.core_properties.title = spec.title
    doc.core_properties.author = "Meeting Scribe"

    def add_hyperlink(paragraph, url: str, text: str):
        part = paragraph.part
        r_id = part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        new_run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        rpr.append(color)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
        new_run.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        new_run.append(t)
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)

    for kind, payload in spec.blocks:
        if kind == "h1":
            doc.add_heading(payload, level=0)
        elif kind == "h2":
            doc.add_heading(payload, level=1)
        elif kind == "h3":
            doc.add_heading(payload, level=2)
        elif kind == "p":
            doc.add_paragraph(payload)
        elif kind == "note":
            p = doc.add_paragraph()
            run = p.add_run(payload)
            run.italic = True
        elif kind == "meta":
            table = doc.add_table(rows=0, cols=2)
            table.style = "Light List Accent 1"
            for label, value in payload:
                row = table.add_row().cells
                row[0].text = str(label)
                row[1].text = str(value)
                for par in row[0].paragraphs:
                    for r in par.runs:
                        r.bold = True
        elif kind == "bullets":
            for item in payload:
                doc.add_paragraph(str(item), style="List Bullet")
        elif kind == "numbered":
            for item in payload:
                doc.add_paragraph(str(item), style="List Number")
        elif kind == "speaker":
            speaker, text, ts = payload
            p = doc.add_paragraph()
            run = p.add_run(f"{speaker}" + (f"  [{ts}]" if ts else "") + ":  ")
            run.bold = True
            p.add_run(text)
        elif kind == "table":
            headers, rows = payload
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Light Grid Accent 1"
            for i, h in enumerate(headers):
                c = table.rows[0].cells[i]
                c.text = str(h)
                for par in c.paragraphs:
                    for r in par.runs:
                        r.bold = True
            for row in rows:
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = str(val)
        elif kind == "link":
            text, target = payload
            p = doc.add_paragraph()
            add_hyperlink(p, target, text)
        elif kind == "hr":
            doc.add_paragraph("―" * 30)
        elif kind == "pagebreak":
            doc.add_page_break()

    doc.save(str(path))


# ======================================================================== ODT
def _odt(spec: DocSpec, path: Path) -> None:
    from odf.opendocument import OpenDocumentText
    from odf.style import (Style, TextProperties, ParagraphProperties,
                           TableColumnProperties, TableCellProperties)
    from odf.text import H, P, Span, A, List as OdfList, ListItem
    from odf.table import Table, TableColumn, TableRow, TableCell

    doc = OpenDocumentText()

    bold = Style(name="Bold", family="text")
    bold.addElement(TextProperties(fontweight="bold"))
    doc.styles.addElement(bold)

    italic = Style(name="Italic", family="text")
    italic.addElement(TextProperties(fontstyle="italic"))
    doc.styles.addElement(italic)

    cell_style = Style(name="Cell", family="table-cell")
    cell_style.addElement(TableCellProperties(border="0.5pt solid #999999",
                                              padding="0.05in"))
    doc.automaticstyles.addElement(cell_style)

    def para(text: str, style_name: str | None = None):
        p = P()
        if style_name:
            p.addElement(Span(stylename=style_name, text=text))
        else:
            p.addText(text)
        doc.text.addElement(p)

    def heading(text: str, level: int):
        doc.text.addElement(H(outlinelevel=level, text=text))

    _tbl_seq = [0]

    def _tname(prefix: str) -> str:
        _tbl_seq[0] += 1
        return f"{prefix}-{_tbl_seq[0]}"

    for kind, payload in spec.blocks:
        if kind == "h1":
            heading(payload, 1)
        elif kind == "h2":
            heading(payload, 2)
        elif kind == "h3":
            heading(payload, 3)
        elif kind == "p":
            para(payload)
        elif kind == "note":
            para(payload, "Italic")
        elif kind == "meta":
            tbl = Table(name=_tname("meta"))
            tbl.addElement(TableColumn())
            tbl.addElement(TableColumn())
            for label, value in payload:
                tr = TableRow()
                c1 = TableCell(stylename="Cell")
                p1 = P()
                p1.addElement(Span(stylename="Bold", text=str(label)))
                c1.addElement(p1)
                c2 = TableCell(stylename="Cell")
                c2.addElement(P(text=str(value)))
                tr.addElement(c1)
                tr.addElement(c2)
                tbl.addElement(tr)
            doc.text.addElement(tbl)
        elif kind in ("bullets", "numbered"):
            lst = OdfList()
            for item in payload:
                li = ListItem()
                li.addElement(P(text=str(item)))
                lst.addElement(li)
            doc.text.addElement(lst)
        elif kind == "speaker":
            speaker, text, ts = payload
            p = P()
            tag = f"{speaker}" + (f"  [{ts}]" if ts else "") + ":  "
            p.addElement(Span(stylename="Bold", text=tag))
            p.addText(text)
            doc.text.addElement(p)
        elif kind == "table":
            headers, rows = payload
            tbl = Table(name=_tname("t"))
            for _ in headers:
                tbl.addElement(TableColumn())
            hr = TableRow()
            for h in headers:
                c = TableCell(stylename="Cell")
                ph = P()
                ph.addElement(Span(stylename="Bold", text=str(h)))
                c.addElement(ph)
                hr.addElement(c)
            tbl.addElement(hr)
            for row in rows:
                tr = TableRow()
                for val in row:
                    c = TableCell(stylename="Cell")
                    c.addElement(P(text=str(val)))
                    tr.addElement(c)
                tbl.addElement(tr)
            doc.text.addElement(tbl)
        elif kind == "link":
            text, target = payload
            p = P()
            a = A(href=target, type="simple")
            a.addText(text)
            p.addElement(a)
            doc.text.addElement(p)
        elif kind == "hr":
            para("―" * 30)
        elif kind == "pagebreak":
            para("")

    doc.save(str(path))


# ======================================================================== PDF
def _pdf(spec: DocSpec, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, ListFlowable, ListItem, PageBreak,
                                    HRFlowable)

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Speaker", parent=styles["BodyText"],
                              spaceBefore=4, spaceAfter=4))
    styles.add(ParagraphStyle(name="NoteStyle", parent=styles["BodyText"],
                              textColor=colors.HexColor("#555555"),
                              leftIndent=12, fontName="Helvetica-Oblique"))

    def esc(s: str) -> str:
        return (str(s).replace("&", "&amp;").replace("<", "&lt;")
                .replace(">", "&gt;"))

    story: list = []
    for kind, payload in spec.blocks:
        if kind == "h1":
            story.append(Paragraph(esc(payload), styles["Title"]))
            story.append(Spacer(1, 8))
        elif kind == "h2":
            story.append(Spacer(1, 6))
            story.append(Paragraph(esc(payload), styles["Heading2"]))
        elif kind == "h3":
            story.append(Paragraph(esc(payload), styles["Heading3"]))
        elif kind == "p":
            story.append(Paragraph(esc(payload), styles["BodyText"]))
        elif kind == "note":
            story.append(Paragraph(esc(payload), styles["NoteStyle"]))
        elif kind == "meta":
            data = [[Paragraph(f"<b>{esc(k)}</b>", styles["BodyText"]),
                     Paragraph(esc(v), styles["BodyText"])] for k, v in payload]
            t = Table(data, colWidths=[1.6 * inch, 4.6 * inch])
            t.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#bbbbbb")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#dddddd")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f2f5fa")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))
        elif kind in ("bullets", "numbered"):
            bt = "1" if kind == "numbered" else "bullet"
            story.append(ListFlowable(
                [ListItem(Paragraph(esc(x), styles["BodyText"]), leftIndent=14)
                 for x in payload],
                bulletType=bt, start="1" if kind == "numbered" else None))
            story.append(Spacer(1, 4))
        elif kind == "speaker":
            speaker, text, ts = payload
            tag = f"<b>{esc(speaker)}</b>" + (f" <font size=8 color='#888'>[{esc(ts)}]</font>" if ts else "")
            story.append(Paragraph(f"{tag}: {esc(text)}", styles["Speaker"]))
        elif kind == "table":
            headers, rows = payload
            data = [[Paragraph(f"<b>{esc(h)}</b>", styles["BodyText"]) for h in headers]]
            for row in rows:
                data.append([Paragraph(esc(v), styles["BodyText"]) for v in row])
            ncol = len(headers)
            width = 6.3 * inch
            if ncol >= 5:
                cw = [0.35 * inch, 2.7 * inch, 1.1 * inch, 1.1 * inch, 1.05 * inch]
            elif ncol == 3:
                cw = [1.3 * inch, 3.2 * inch, 1.8 * inch]
            else:
                cw = [width / ncol] * ncol
            t = Table(data, colWidths=cw, repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8edf5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))
        elif kind == "link":
            text, target = payload
            story.append(Paragraph(
                f'<link href="{esc(target)}"><font color="#0563C1">{esc(text)}</font></link>',
                styles["BodyText"]))
        elif kind == "hr":
            story.append(Spacer(1, 6))
            story.append(HRFlowable(width="100%", thickness=0.6,
                                    color=colors.HexColor("#cccccc")))
            story.append(Spacer(1, 6))
        elif kind == "pagebreak":
            story.append(PageBreak())

    doc = SimpleDocTemplate(str(path), pagesize=LETTER,
                            title=spec.title, author="Meeting Scribe",
                            leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                            topMargin=0.9 * inch, bottomMargin=0.9 * inch)
    doc.build(story)


def attach_file_to_pdf(pdf_path: Path, media_path: Path) -> None:
    """Embed the source recording inside the PDF as a file attachment."""
    try:
        from pypdf import PdfReader, PdfWriter

        reader = PdfReader(str(pdf_path))
        writer = PdfWriter()
        writer.append(reader)
        with open(media_path, "rb") as fh:
            writer.add_attachment(media_path.name, fh.read())
        tmp = pdf_path.with_suffix(".tmp.pdf")
        with open(tmp, "wb") as fh:
            writer.write(fh)
        tmp.replace(pdf_path)
        log.info("Embedded %s into %s", media_path.name, pdf_path.name)
    except Exception as e:  # pragma: no cover
        log.warning("Could not embed audio into PDF (%s); leaving a link instead.", e)
