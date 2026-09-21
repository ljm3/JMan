"""Pipeline behaviour with fake carriers: fallback, missing keys, cache, POD, Claude stub, Google request building."""
import json
from datetime import datetime
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from tracking_checker import carriers, config
from tracking_checker import models as M
from tracking_checker.carriers.base import NotConfigured, NotFound
from tracking_checker.pipeline import RunRequest, run
from tracking_checker.selftest import fedex12, mod10, ups_number

UPS_N = ups_number("AB1234567890123")
FEDEX_N = fedex12("27777777777")
USPS92 = mod10("92" + "1" * 19)          # USPS first, FedEx second


def book(tmp_path, numbers):
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Order", "Carrier", "Tracking Number"])
    for i, n in enumerate(numbers):
        ws.append([f"SO-{i}", "", n])
    p = tmp_path / "orders.xlsx"
    wb.save(p)
    return p


class Fake:
    batch_size = 1

    def __init__(self, name, known, calls):
        self.name, self.known, self.calls = name, known, calls

    def track_many(self, nums):
        out = {}
        for n in nums:
            self.calls.append((self.name, n))
            if n in self.known:
                r = M.TrackingResult(n, self.name, status=M.DELIVERED, delivered_at=datetime(2026, 9, 2, 9, 30),
                                     events=[M.TrackingEvent(datetime(2026, 9, 2, 9, 30), "Delivered", "Dallas, TX")])
                r.pod = M.ProofOfDelivery(signed_by="SMITH", signature_image=b"GIF89a....")
                out[n] = r
            else:
                out[n] = NotFound(f"{self.name}: not found")
        return out

    def fetch_pod(self, r):
        r.pod.document = b"%PDF-1.4 fake"
        r.pod.document_label = f"{self.name} POD"


def fake_build(known, configured=("UPS", "FedEx", "USPS"), calls=None):
    calls = calls if calls is not None else []

    def build(carrier, settings, demo=False, cancel=None, **_):
        if carrier not in configured:
            raise NotConfigured(f"No {carrier} API credentials - add them in Settings.")
        return Fake(carrier, known.get(carrier, set()), calls)
    return build, calls


def settings():
    s = config.Settings()
    s.claude_enabled = False
    return s


def test_fallback_to_second_carrier_and_pod(tmp_path, monkeypatch):
    p = book(tmp_path, [UPS_N, USPS92])
    build, calls = fake_build({"UPS": {UPS_N}, "FedEx": {USPS92}})
    monkeypatch.setattr(carriers, "build", build)
    s = run(RunRequest(str(p), "Orders", 3), settings())
    assert ("USPS", USPS92) in calls and ("FedEx", USPS92) in calls
    ws = load_workbook(p)["Tracking Status"]
    hdr = [c.value for c in ws[1]]
    rows = {ws.cell(r, 1).value: {h: ws.cell(r, i + 1) for i, h in enumerate(hdr)} for r in range(2, ws.max_row + 1)}
    assert rows[USPS92]["Carrier"].value == "FedEx" and "found by FedEx" in rows[USPS92]["Identified By"].value
    pod = rows[UPS_N]["Proof of Delivery"]
    assert pod.value == "View POD" and pod.hyperlink.target.endswith("_POD.pdf")
    assert (p.parent / pod.hyperlink.target).exists()
    assert "signature image" not in rows[UPS_N]["POD Type"].value   # carrier doc wins
    assert s.pod_files >= 2


def test_missing_keys_are_reported_not_fatal(tmp_path, monkeypatch):
    p = book(tmp_path, [UPS_N, FEDEX_N])
    build, _ = fake_build({"UPS": {UPS_N}}, configured=("UPS",))
    monkeypatch.setattr(carriers, "build", build)
    s = run(RunRequest(str(p), "Orders", 3), settings())
    ws = load_workbook(p)["Tracking Status"]
    hdr = [c.value for c in ws[1]]
    fed = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == FEDEX_N)
    assert ws.cell(fed, hdr.index("Status") + 1).value == M.ERROR
    assert "No FedEx API credentials" in ws.cell(fed, hdr.index("Error / Notes") + 1).value
    assert any("FedEx" in n for n in s.notes)


def test_delivered_results_are_cached_and_reused(tmp_path, monkeypatch):
    p = book(tmp_path, [UPS_N])
    build, calls = fake_build({"UPS": {UPS_N}})
    monkeypatch.setattr(carriers, "build", build)
    run(RunRequest(str(p), "Orders", 3), settings())
    # the fake returns source "Live API" by default, so it is cached
    assert len(calls) == 1
    run(RunRequest(str(p), "Orders", 3), settings())
    assert len(calls) == 1, "second run should re-use the delivered result"
    ws = load_workbook(p)["Tracking Status"]
    hdr = [c.value for c in ws[1]]
    assert ws.cell(2, hdr.index("Data Source") + 1).value.startswith("Cached")


def test_rerun_replaces_results_tab(tmp_path, monkeypatch):
    p = book(tmp_path, [UPS_N])
    monkeypatch.setattr(carriers, "build", fake_build({"UPS": {UPS_N}})[0])
    run(RunRequest(str(p), "Orders", 3, reuse_delivered=False), settings())
    run(RunRequest(str(p), "Orders", 3, reuse_delivered=False), settings())
    wb = load_workbook(p)
    assert wb.sheetnames == ["Orders", "Tracking Status"]
    assert wb["Orders"]["C2"].hyperlink.location == "'Tracking Status'!A2"


def test_refuses_results_tab_as_source(tmp_path):
    p = book(tmp_path, [UPS_N])
    with pytest.raises(ValueError):
        run(RunRequest(str(p), "Tracking Status", 3, status_tab="Tracking Status"), settings())


# ---------------------------------------------------------------- Claude path (stubbed client)
class StubMessages:
    def __init__(self, replies):
        self.replies, self.calls = replies, []

    def create(self, **kw):
        self.calls.append(kw)
        text = self.replies.pop(0)
        return SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text=text)])


def test_claude_check_merges_values(monkeypatch):
    from tracking_checker import checks
    from tracking_checker.checks import claude_check

    plan = {"interpretation": "flag signer mismatch", "columns": [{"name": "Signer Is Smith", "instruction": "..."}],
            "cannot_do": ["invoice totals"]}
    judged = {"rows": [{"tracking_number": UPS_N, "values": [{"column": "Signer Is Smith", "value": "Yes"}],
                        "matches": True, "note": "signed by SMITH"}]}
    stub = StubMessages([json.dumps(plan), json.dumps(judged)])

    class FakeChecker(claude_check.ClaudeChecker):
        def __init__(self, key, model):
            self.model, self._fallbacks_ok = model, True
            import anthropic
            self._anthropic = anthropic
            self.client = SimpleNamespace(beta=SimpleNamespace(messages=stub), messages=stub)

    monkeypatch.setattr(claude_check, "ClaudeChecker", FakeChecker)
    monkeypatch.setattr(config, "get_secret", lambda name: "sk-test" if name == "anthropic_api_key" else "")
    s = config.Settings()
    r1 = M.TrackingResult(UPS_N, M.UPS, status=M.DELIVERED)
    r2 = M.TrackingResult(FEDEX_N, M.FEDEX, status=M.IN_TRANSIT)
    out = checks.run("is the signer Smith?", [r1, r2], s)
    assert out.engine.startswith("Claude") and out.columns == ["Signer Is Smith", "Matches Your Check", "Check Notes"]
    assert r1.extra["Signer Is Smith"] == "Yes" and r1.extra["Matches Your Check"] == "Yes"
    assert r2.extra["Matches Your Check"] == "Unknown"
    assert any("invoice" in n for n in out.notes)
    first = stub.calls[0]
    assert first["model"] == "claude-opus-5" and first["fallbacks"] == "default"
    assert first["thinking"] == {"type": "adaptive"} and first["output_config"]["format"]["type"] == "json_schema"


def test_claude_failure_falls_back_to_rules(monkeypatch):
    from tracking_checker import checks
    from tracking_checker.checks import claude_check

    def boom(*a, **k):
        raise claude_check.ClaudeCheckError("network down")

    monkeypatch.setattr(claude_check.ClaudeChecker, "__init__", lambda self, k, m: None)
    monkeypatch.setattr(claude_check.ClaudeChecker, "plan", boom)
    monkeypatch.setattr(config, "get_secret", lambda name: "sk-test" if name == "anthropic_api_key" else "")
    r = M.TrackingResult(UPS_N, M.UPS, status=M.IN_TRANSIT)
    out = checks.run("not delivered yet", [r], config.Settings())
    assert out.engine == "Built-in rules" and r.extra["Not Yet Delivered"] == "Yes"
    assert "network down" in out.notes[0]


# ---------------------------------------------------------------- Google Sheets request building
def test_gsheets_links_and_formats():
    from tracking_checker.report import OutCell
    from tracking_checker.sheets.gsheets import GoogleWorkbook

    class WS:
        def __init__(self, title, gid):
            self.title, self.id, self.updated = title, gid, None

        def clear(self):
            pass

        def resize(self, **k):
            pass

        def update(self, values, range_name, value_input_option):
            self.updated = (values, value_input_option)

    tabs = {"Orders": WS("Orders", 11), "Tracking Status": WS("Tracking Status", 22)}
    sent = []
    wb = GoogleWorkbook.__new__(GoogleWorkbook)
    wb.sh = SimpleNamespace(worksheet=lambda t: tabs[t], batch_update=lambda body: sent.append(body))
    wb._pending, wb.url = [], "https://docs.google.com/spreadsheets/d/x/edit"
    rows = [[OutCell(UPS_N, goto=("Orders", 2, 3)), OutCell("Delivered", fill="C6EFCE"),
             OutCell("View POD", url="https://drive.google.com/file/d/abc/view")]]
    wb.write_status_tab("Tracking Status", ["Tracking Number", "Status", "Proof of Delivery"], rows, {},
                        {"Status": "note text"})
    wb.link_source_cells("Orders", [(2, 3, 2)], "Tracking Status")
    wb.save()
    values, mode = tabs["Tracking Status"].updated
    assert mode == "RAW" and values[1][0] == UPS_N
    reqs = [r for body in sent for r in body["requests"]]
    links = [r["repeatCell"]["cell"]["userEnteredFormat"]["textFormat"]["link"]["uri"] for r in reqs
             if "repeatCell" in r and "textFormat" in r["repeatCell"]["cell"].get("userEnteredFormat", {})
             and "link" in r["repeatCell"]["cell"]["userEnteredFormat"]["textFormat"]]
    assert "#gid=11&range=C2" in links                       # status row -> original cell
    assert "#gid=22&range=A2" in links                       # original cell -> status row
    assert "https://drive.google.com/file/d/abc/view" in links
    assert any(r.get("repeatCell", {}).get("cell", {}).get("note") == "note text" for r in reqs)
