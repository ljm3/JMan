"""Website-mode parsing and method selection.

ups_web_getstatus.json is a trimmed copy of what www.ups.com/track actually returned (2026-09-19) for
UPS's public example number 1Z999AA10123456784.
"""
import json
import os
from pathlib import Path

import pytest

from tracking_checker import carriers, config
from tracking_checker import models as M
from tracking_checker.carriers.base import NotFound
from tracking_checker.carriers.web.fedex_web import parse_fedex_web
from tracking_checker.carriers.web.ups_web import parse_ups_web
from tracking_checker.carriers.web.usps_web import parse_usps_web

FIX = Path(__file__).parent / "fixtures"


def test_ups_web_real_response():
    data = json.loads((FIX / "ups_web_getstatus.json").read_text())
    r = parse_ups_web("1Z999AA10123456784", data)
    assert r.status == M.DELIVERED and r.source == "Carrier website"
    assert r.service.startswith("UPS Express")
    assert r.delivered_at.strftime("%Y-%m-%d %H:%M") == "2025-12-05 10:29"
    assert r.pod.signed_by == "TAYLOR" and r.pod.left_at == ""          # "Other" is dropped
    assert r.destination == "Longview, TX" and len(r.events) == 17
    assert r.events[0].description == "Delivered" and r.events[0].location == "Longview, TX"


def test_ups_web_error_is_not_found():
    with pytest.raises(NotFound):
        parse_ups_web("1Z", {"statusCode": "200", "trackDetails": [{"errorCode": "504", "errorText": "Not found"}]})


def test_ups_web_signature_when_signed_in():
    data = json.loads((FIX / "ups_web_getstatus.json").read_text())
    data["trackDetails"][0]["signatureImage"] = "R0lGODlhAQABAAAAACw="
    r = parse_ups_web("1Z999AA10123456784", data)
    assert r.pod.signature_image.startswith(b"GIF")


USPS_DELIVERED = {
    "status": "Delivered",
    "banner": "",
    "body": "Tracking Number:\n9400111899223100000007\nProduct Information\nPostal Product:\nUSPS Ground Advantage\n",
    "steps": [
        {"status": "Delivered", "detail": "Delivered, In/At Mailbox", "location": "DALLAS, TX 75261",
         "date": "September 4, 2026, 1:05 pm", "lines": []},
        {"status": "", "detail": "Out for Delivery", "location": "DALLAS, TX 75261",
         "date": "September 4, 2026, 7:10 am", "lines": []},
        {"status": "", "detail": "", "location": "", "date": "",
         "lines": ["USPS in possession of item", "PHOENIX, AZ 85034", "August 31, 2026, 4:00 pm"]},
    ],
}


def test_usps_web_delivered():
    r = parse_usps_web("9400111899223100000007", USPS_DELIVERED)
    assert r.status == M.DELIVERED and r.service == "USPS Ground Advantage"
    assert r.delivered_at.strftime("%Y-%m-%d %H:%M") == "2026-09-04 13:05"
    assert r.pod.left_at == "In/At Mailbox" and r.destination == "Dallas, TX 75261"
    assert r.origin == "Phoenix, AZ 85034" and len(r.events) == 3      # third step parsed from its text lines


def test_usps_web_not_available_is_not_found():
    body = "Tracking Number:\n9400\nTracking Not Available\nTracking is not available for this item."
    with pytest.raises(NotFound):
        parse_usps_web("9400", {"steps": [], "body": body})


def test_fedex_web_uses_track_api_shape():
    data = {"output": {"completeTrackResults": [{"trackingNumber": "123456789012", "trackResults": [{
        "latestStatusDetail": {"code": "IT", "derivedCode": "IT", "statusByLocale": "In transit",
                               "description": "In transit"},
        "scanEvents": [{"date": "2026-09-02T08:00:00-05:00", "eventDescription": "Departed FedEx hub",
                        "scanLocation": {"city": "MEMPHIS", "stateOrProvinceCode": "TN"}}]}]}]}}
    r = parse_fedex_web("123456789012", data)
    assert r.status == M.IN_TRANSIT and r.source == "Carrier website"


def test_lookup_method_resolution(monkeypatch):
    s = config.Settings()
    assert config.lookup_method("UPS", s) == config.WEBSITE                 # auto, no keys
    monkeypatch.setattr(config, "carrier_configured", lambda c: c == "FedEx")
    monkeypatch.setattr(config, "carrier_credentials", lambda c: ("id", "secret") if c == "FedEx" else ("", ""))
    assert config.lookup_method("FedEx", s) == config.API                   # auto, keys set
    s.ups.method = config.API
    assert config.lookup_method("UPS", s) == config.API                     # forced
    assert type(carriers.build("USPS", s, browser=object())).__name__ == "USPSWeb"
    assert type(carriers.build("FedEx", s, browser=object())).__name__ == "FedExCarrier"


@pytest.mark.skipif(os.environ.get("TC_LIVE_WEB") != "1", reason="set TC_LIVE_WEB=1 to hit ups.com")
def test_live_ups_website(tmp_path, monkeypatch):
    from openpyxl import Workbook, load_workbook

    from tracking_checker.pipeline import RunRequest, run
    monkeypatch.delenv("TC_BROWSER_PROFILE", raising=False)
    wb = Workbook()
    wb.active.append(["Tracking"])
    wb.active.append(["1Z999AA10123456784"])
    p = tmp_path / "live.xlsx"
    wb.save(p)
    s = config.Settings()
    s.claude_enabled = False
    run(RunRequest(str(p), "Sheet", 1), s)
    st = load_workbook(p)["Tracking Status"]
    hdr = [c.value for c in st[1]]
    assert st.cell(2, hdr.index("Status") + 1).value == "Delivered"
    assert st.cell(2, hdr.index("Data Source") + 1).value == "Carrier website"
