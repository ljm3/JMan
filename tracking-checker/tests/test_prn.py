"""PRN tab: PRN/ZIP detection, UPS pickup parsing, and the run writing results + links + columns."""
import json
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from tracking_checker import config
from tracking_checker.carriers.base import NotFound
from tracking_checker.pickup import (AUTO_ZIP, COMPLETED, INCOMPLETE, NO_ZIP, NOT_FOUND, NOT_SUPPORTED, PROCESSING,
                                     is_ups_prn, status_from_ups, zip_from_value)
from tracking_checker.pickup.pipeline import PrnRequest, plan_columns, run_prn
from tracking_checker.pickup.sheet import analyse_prn_columns, find_prns
from tracking_checker.pickup.ups_web import parse_ups_pickup
from tracking_checker.selftest import ups_number

FIX = json.loads((Path(__file__).parent / "fixtures" / "ups_pickup_getdetails.json").read_text())


# ---------------------------------------------------------------- ZIP codes
@pytest.mark.parametrize("value,zip_col,expected", [
    ("99 Park Ave, New York, NY 10016", False, ("10016", "US")),
    ("31-13 30th Ave, Astoria NY 11102-1234", False, ("11102", "US")),
    ("12345 Main St, Dallas, TX 75201", False, ("75201", "US")),       # house number looks like a ZIP
    ("PO Box 55555, Hoboken NJ 07030", False, ("07030", "US")),
    ("1402 Avenue J\nBrooklyn 11230", False, ("11230", "US")),        # no state code: last 5-digit group
    ("10016 Main Street", False, None),                                # only the house number
    ("07030", False, ("07030", "US")),
    ("10016-4411", False, ("10016", "US")),
    (7030, True, ("07030", "US")),                                     # Excel dropped the leading zero
    (7030, False, None),
    (100164411, False, ("10016", "US")),
    ("100 King St W, Toronto, ON M5X 1A9", False, ("M5X 1A9", "CA")),
    ("", False, None),
    (None, False, None),
])
def test_zip_from_value(value, zip_col, expected):
    assert zip_from_value(value, zip_column=zip_col) == expected


def test_prn_pattern():
    assert is_ups_prn("2900000AA11") and is_ups_prn("292222CC33D")
    assert not is_ups_prn("2981220CD8")          # 10 characters
    assert not is_ups_prn("ABC1220CD89")         # starts with a letter
    assert not is_ups_prn("12125550100")         # 11-digit phone number
    assert not is_ups_prn(ups_number("AB1234567890123"))


def test_status_mapping():
    assert status_from_ups("003", "Your pickup request has been successfully completed.") == COMPLETED
    assert status_from_ups("002", "Your pickup request has been received") == PROCESSING
    assert status_from_ups("004", "UPS attempted to pickup your shipment(s) but was unsuccessful.") == INCOMPLETE
    assert status_from_ups("009", "Your pickup was cancelled") == "Cancelled"


# ---------------------------------------------------------------- parsing (real responses, anonymised)
def test_parse_completed():
    r = parse_ups_pickup("2900000AA11", "10016", "US", FIX["2900000AA11"])
    assert r.status == COMPLETED and r.status_code == "003"
    assert r.status_time == datetime(2026, 9, 21, 13, 51, 47)
    assert r.pickup_date == datetime(2026, 9, 21) and r.window() == "9:00 AM - 4:00 PM"
    assert r.pieces == [("UPS Ground", 4)] and r.piece_count == 4
    assert r.weight == "70 lbs"
    assert r.address == "10 Sample Ave, New York, NY 10016"
    assert r.phone == "(555) 010-0000" and r.pickup_point == "Front Door" and r.residential is False
    assert (r.base_charge, r.fuel_surcharge, r.other_surcharges, r.total_charge) == (9.65, 2.85, 0.0, 12.5)
    assert r.tracking_numbers == []


def test_parse_processing_and_incomplete():
    assert parse_ups_pickup("2911111BB22", "11102", "US", FIX["2911111BB22"]).status == PROCESSING
    r = parse_ups_pickup("292222CC33D", "11230", "US", FIX["292222CC33D"])
    assert r.status == INCOMPLETE and "unsuccessful" in r.status_detail


def test_parse_error_and_empty():
    with pytest.raises(NotFound):
        parse_ups_pickup("2900000AA11", "10016", "US", {})
    with pytest.raises(NotFound):
        parse_ups_pickup("2900000AA11", "10016", "US", {"errorCode": "9510110", "errorMessage": "Invalid PRN"})


# ---------------------------------------------------------------- sheet detection
def sheet_rows():
    return [
        ["Site", "Carrier", "Pickup #", "Ship From Address", "Zip", "Notes"],
        ["NYC", "UPS", "2900000AA11", "10 Sample Ave, New York, NY 10016", None, "call first"],
        ["Astoria", "UPS", "2911111BB22", "20-20 Example Ave, Astoria NY", "11102", None],
        ["Brooklyn", "UPS", "292222CC33D", "30 Avenue Q, Brooklyn, NY 11230", 11230, None],
        ["Dallas", "FedEx", "CLTA123456", "1 Main St, Dallas, TX 75201", 75201, None],
        ["Nowhere", "UPS", "2933333DD44", "no address given", None, None],
        ["NYC again", "UPS", "2900000AA11", "10 Sample Ave, New York, NY 10016", "10016", None],
    ]


def test_analyse_and_find():
    grid = sheet_rows()
    hdr, prn_cols, zip_cols = analyse_prn_columns(grid)
    assert hdr == 0 and prn_cols[0].index == 3
    assert [(z.index, z.kind) for z in zip_cols][:2] == [(5, "zip"), (4, "address")]
    found = find_prns(grid, 3, hdr, AUTO_ZIP, zip_cols)
    p = found.pickups
    assert set(p) == {"2900000AA11", "2911111BB22", "292222CC33D", "CLTA123456", "2933333DD44"}
    assert p["2900000AA11"].zip_code == "10016" and "address" in p["2900000AA11"].zip_source   # Zip cell empty
    assert p["2900000AA11"].occurrences == [(2, 3), (7, 3)]
    assert p["2911111BB22"].zip_code == "11102" and "ZIP column" in p["2911111BB22"].zip_source
    assert p["CLTA123456"].carrier == "FedEx" and not p["CLTA123456"].supported
    assert p["2933333DD44"].zip_code == ""
    assert "4 UPS" in found.summary() and "1 other" in found.summary()


def test_plan_columns_reuses_ours():
    grid = [["A", "B", None, "PRN Status"], [1, 2, None, "x"]]
    assert plan_columns(grid, 0, ["PRN Status", "PRN Pieces"]) == {"PRN Status": 4, "PRN Pieces": 5}
    assert plan_columns([["A", "B"], [1, None]], 0, ["PRN Status"]) == {"PRN Status": 3}


# ---------------------------------------------------------------- full run (Excel)
class FakeUPS:
    def __init__(self):
        self.calls = []

    def check(self, prn, zip_code, country="US"):
        self.calls.append((prn, zip_code))
        if prn not in FIX:
            raise NotFound(f"UPS has no pickup {prn} for ZIP {zip_code}")
        return parse_ups_pickup(prn, zip_code, country, FIX[prn])


def make_book(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pickups"
    for r in sheet_rows():
        ws.append(r)
    ws["F1"].style = "Accent1"
    wb.create_sheet("Other")["A1"] = "unrelated"
    p = tmp_path / "pickups.xlsx"
    wb.save(p)
    return p


def test_run_writes_results_links_and_columns(tmp_path):
    path = make_book(tmp_path)
    fake = FakeUPS()
    req = PrnRequest(location=str(path), sheet="Pickups", column=3)
    s = run_prn(req, config.Settings(), client=fake)
    # one lookup per unique UPS PRN that has a ZIP (no FedEx, no ZIP-less, no duplicate)
    assert fake.calls == [("2900000AA11", "10016"), ("2911111BB22", "11102"), ("292222CC33D", "11230")]
    assert s.total == 5 and s.by_status[COMPLETED] == 1 and s.by_status[NOT_SUPPORTED] == 1
    assert s.by_status[NO_ZIP] == 1

    wb = load_workbook(path)
    rs = wb["PRN Status"]
    heads = [c.value for c in rs[1]]
    col = {h: i + 1 for i, h in enumerate(heads)}
    rows = {rs.cell(r, 1).value: r for r in range(2, rs.max_row + 1)}
    r = rows["2900000AA11"]
    assert rs.cell(r, col["Status"]).value == COMPLETED
    assert rs.cell(r, col["Pieces"]).value == 4
    assert rs.cell(r, col["Pickup Address"]).value == "10 Sample Ave, New York, NY 10016"
    assert "None listed by UPS (4" in rs.cell(r, col["Tracking Numbers"]).value
    assert rs.cell(r, col["Source Location"]).value == "'Pickups'!C2; C7"
    # results tab -> original row, and original PRN cell -> results row
    assert rs.cell(r, 1).hyperlink.location == "'Pickups'!C2"
    src = wb["Pickups"]
    assert src["C2"].hyperlink.location == f"'PRN Status'!A{r}"
    assert src["C7"].hyperlink.location == f"'PRN Status'!A{r}"
    assert rs.cell(rows["CLTA123456"], col["Status"]).value == NOT_SUPPORTED
    assert rs.cell(rows["2933333DD44"], col["Status"]).value == NO_ZIP

    # write-back columns: first empty columns right of the data (F is the last used -> G, H)
    assert src["G1"].value == "PRN Status" and src["H1"].value == "PRN Pieces"
    assert src["G2"].value == COMPLETED and src["H2"].value == 4
    assert src["G3"].value == PROCESSING and src["G4"].value == INCOMPLETE
    assert src["G5"].value == NOT_SUPPORTED and src["G6"].value == NO_ZIP
    assert src["G2"].fill.fgColor.rgb.endswith("C6EFCE")
    assert src["G1"].fill.fgColor.theme == src["F1"].fill.fgColor.theme == 4                       # header styled like its neighbour
    assert src["A2"].value == "NYC" and src["F2"].value == "call first"   # original data untouched


def test_rerun_reuses_columns(tmp_path):
    path = make_book(tmp_path)
    req = PrnRequest(location=str(path), sheet="Pickups", column=3)
    run_prn(req, config.Settings(), client=FakeUPS())
    run_prn(req, config.Settings(), client=FakeUPS())
    src = load_workbook(path)["Pickups"]
    heads = [c.value for c in src[1]]
    assert heads.count("PRN Status") == 1 and heads.count("PRN Pieces") == 1
    assert len(heads) == 8


def test_no_write_back_and_not_found(tmp_path):
    path = make_book(tmp_path)

    class Missing(FakeUPS):
        def check(self, prn, zip_code, country="US"):
            raise NotFound("UPS has no pickup")

    s = run_prn(PrnRequest(location=str(path), sheet="Pickups", column=3, write_back=False), config.Settings(),
                client=Missing())
    assert s.by_status[NOT_FOUND] == 3 and not s.columns_added
    assert load_workbook(path)["Pickups"].max_column == 6


def test_demo_mode(tmp_path):
    path = make_book(tmp_path)
    s = run_prn(PrnRequest(location=str(path), sheet="Pickups", column=3, demo=True), config.Settings())
    assert s.total == 5 and any("DEMO" in n for n in s.notes)
    rs = load_workbook(path)["PRN Status"]
    assert rs.cell(2, [c.value for c in rs[1]].index("Data Source") + 1).value.startswith("DEMO")


def test_rejects_results_tab():
    with pytest.raises(ValueError):
        run_prn(PrnRequest(location="x.xlsx", sheet="PRN Status", column=3), config.Settings(), client=FakeUPS())


# ---------------------------------------------------------------- Google Sheets request building
def test_gsheets_write_columns():
    from tracking_checker.report import OutCell
    from tracking_checker.sheets.gsheets import GoogleWorkbook

    calls = SimpleNamespace(values=None, added=0)
    ws = SimpleNamespace(id=7, col_count=6, row_count=10,
                         add_cols=lambda n: setattr(calls, "added", n), add_rows=lambda n: None,
                         batch_update=lambda data, value_input_option: setattr(calls, "values", data))
    gw = GoogleWorkbook.__new__(GoogleWorkbook)
    gw.sh = SimpleNamespace(worksheet=lambda name: ws)
    gw._pending = []
    gw.write_columns("Pickups", 1, [(7, "PRN Status", {2: OutCell("Completed", fill="C6EFCE")})], 2, 3)
    assert calls.added == 1
    assert calls.values == [{"range": "G1", "values": [["PRN Status"]]},
                            {"range": "G2:G3", "values": [["Completed"], [""]]}]
    fills = [r["repeatCell"]["cell"]["userEnteredFormat"].get("backgroundColor") for r in gw._pending]
    assert fills[1]["green"] == pytest.approx(0xEF / 255) and fills[2] == {"red": 1, "green": 1, "blue": 1}
