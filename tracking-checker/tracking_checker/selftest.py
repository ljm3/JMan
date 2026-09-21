"""Sample workbook generator + end-to-end self-test (demo data, no network)."""
from __future__ import annotations

import random
import tempfile
from datetime import date, timedelta
from pathlib import Path


def ups_number(body15: str) -> str:
    total = 0
    for i, ch in enumerate(body15):
        v = int(ch) if ch.isdigit() else (ord(ch) - ord("A") + 2) % 10
        total += v * (2 if i % 2 else 1)
    return f"1Z{body15}{(10 - total % 10) % 10}"


def fedex12(prefix11: str) -> str:
    w = [1, 3, 7]
    total = sum(int(d) * w[i % 3] for i, d in enumerate(reversed(prefix11)))
    return prefix11 + str((total % 11) % 10)


def mod10(body: str) -> str:
    total = sum(int(d) * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(body)))
    return body + str((10 - total % 10) % 10)


def sample_numbers(seed: int = 7) -> list[tuple[str, str]]:
    rnd = random.Random(seed)
    digits = lambda n: "".join(rnd.choice("0123456789") for _ in range(n))  # noqa: E731
    out = []
    for _ in range(4):
        out.append(("UPS", ups_number("A" + rnd.choice("BCDEFGHJKLMNPRSTVWXY") + digits(13))))
    for _ in range(4):
        out.append(("FedEx", fedex12(rnd.choice("27") + digits(10))))
    for _ in range(3):
        out.append(("USPS", mod10("9400" + digits(17))))
    out.append(("USPS", mod10("7022" + digits(15))))
    return out


def make_sample(path: str | Path) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Shipments"
    ws.append(["Order #", "Customer", "Ship Via", "Tracking #", "Ship Date", "Notes"])
    for c in ws[1]:
        c.font = Font(bold=True)
    rows = sample_numbers()
    customers = ["Acme Corp", "Globex", "Initech", "Umbrella", "Stark Ind.", "Wayne Ent.", "Hooli", "Vandelay"]
    start = date.today() - timedelta(days=14)
    for i, (carrier, num) in enumerate(rows):
        shown = num
        if carrier == "USPS" and i % 2:
            shown = " ".join(num[j:j + 4] for j in range(0, len(num), 4))   # people often paste with spaces
        ws.append([f"SO-{1001 + i}", customers[i % len(customers)], carrier, shown, start + timedelta(days=i % 6), ""])
    ws.append(["SO-2001", "Acme Corp", "UPS", rows[0][1], start, "Second box on the same shipment (duplicate number)"])
    ws.append(["SO-2002", "Globex", "USPS", 9.400111899223e21, start, "Stored as a number - Excel lost digits"])
    for col, w in zip("ABCDEF", (10, 14, 10, 30, 12, 44)):
        ws.column_dimensions[col].width = w
    other = wb.create_sheet("Contacts")
    other.append(["Name", "Phone"])
    other.append(["Jane Doe", "555-201-3344"])
    wb.save(path)
    return str(path)


def selftest() -> int:
    import os

    from openpyxl import load_workbook

    from . import config, report
    from .pipeline import RunRequest, run
    from .sheets import analyse_columns, open_workbook

    tmp = Path(tempfile.mkdtemp(prefix="tracking-checker-selftest-"))
    os.environ.setdefault("TC_WORKSPACE", str(tmp / "workspace"))
    path = Path(make_sample(tmp / "sample.xlsx"))
    checks_failed: list[str] = []

    def check(cond, label):
        print(("  PASS  " if cond else "  FAIL  ") + label)
        if not cond:
            checks_failed.append(label)

    print(f"Self-test workbook: {path}")
    settings = config.Settings()
    settings.claude_enabled = False          # keep the self-test offline and free
    wb = open_workbook(str(path), settings)
    hdr, cols = analyse_columns(wb.read_tab("Shipments"))
    check(hdr == 0, "header row detected")
    check(bool(cols) and cols[0].index == 4, "tracking column D auto-detected")
    req = RunRequest(location=str(path), sheet="Shipments", column=cols[0].index,
                     extra_check="Flag anything not delivered yet. Delivered after 1/1/2020. Signed by smith.",
                     status_tab="Tracking Status", demo=True)
    summary = run(req, settings, lambda m, f=None: None, workbook=wb)
    print(summary.text().replace("\n", "\n    "))

    out = load_workbook(path)
    check("Tracking Status" in out.sheetnames, "results tab created")
    st, src = out["Tracking Status"], out["Shipments"]
    headers = [c.value for c in st[1]]
    base = [h for h, _ in report.BASE_HEADERS]
    check(headers[:len(base)] == base, "all standard headers written")
    check("Matches Your Check" in headers and "Not Yet Delivered" in headers, "extra-check columns added")
    note_cell = st.cell(1, headers.index("Matches Your Check") + 1) if "Matches Your Check" in headers else None
    check(note_cell is not None and note_cell.comment is not None and "smith" in note_cell.comment.text,
          "your request + its interpretation noted on the 'Matches Your Check' header")
    uniq = len(sample_numbers()) + 1          # +1 = the precision-loss cell reported as an error row
    check(st.max_row - 1 == uniq, f"one row per unique number ({st.max_row - 1} == {uniq})")

    ok_links = True
    for r in range(2, src.max_row + 1):
        cell = src.cell(r, 4)
        link = cell.hyperlink
        if link is None or not link.location or "Tracking Status" not in link.location:
            ok_links = False
            print(f"        missing link on Shipments!D{r}")
            continue
        target = int(link.location.split("!A")[1])
        want = str(cell.value).replace(" ", "")
        got = str(st.cell(target, 1).value)
        if got != want:
            ok_links = False
            print(f"        Shipments!D{r} -> row {target} holds {got}, expected {want}")
    check(ok_links, "every source tracking number links to its matching status row")

    back = all(st.cell(r, 1).hyperlink is not None and "Shipments" in (st.cell(r, 1).hyperlink.location or "")
               for r in range(2, st.max_row + 1))
    check(back, "every status row links back to the original row")

    col = {h: i + 1 for i, h in enumerate(headers)}
    delivered = [r for r in range(2, st.max_row + 1) if st.cell(r, col["Status"]).value == "Delivered"]
    pod_ok = all(st.cell(r, col["Proof of Delivery"]).hyperlink is not None and
                 (path.parent / st.cell(r, col["Proof of Delivery"]).hyperlink.target).exists() for r in delivered)
    check(bool(delivered) and pod_ok, f"POD link + file for every delivered row ({len(delivered)})")
    demo_marked = all("DEMO" in str(st.cell(r, col["Data Source"]).value or "") or st.cell(r, col["Status"]).value in
                      ("Error", "Not Found") for r in range(2, st.max_row + 1))
    check(demo_marked, "demo rows are labelled DEMO")
    err_rows = [r for r in range(2, st.max_row + 1) if "lost" in str(st.cell(r, col["Error / Notes"]).value or "")]
    check(len(err_rows) == 1, "number stored as an Excel number is flagged")
    check(any(p.name.startswith("sample_") for p in (config.workspace_dir() / "backups").iterdir()),
          "backup of the original workbook kept")

    print()
    if checks_failed:
        print(f"SELF-TEST FAILED ({len(checks_failed)} check(s)).")
        return 1
    print("SELF-TEST PASSED.")
    return 0
