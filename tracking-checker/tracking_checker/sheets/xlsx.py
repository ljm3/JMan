from __future__ import annotations

import logging
import re
import shutil
import zipfile
from copy import copy
from datetime import date, datetime
from pathlib import Path

from ..config import workspace_dir
from ..report import OutCell, file_link

log = logging.getLogger("tracking_checker")


def quote_tab(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


class XlsxWorkbook:
    kind = "xlsx"

    def __init__(self, path: str | Path):
        import openpyxl

        self.path = Path(path)
        ext = self.path.suffix.lower()
        if ext == ".xls":
            raise ValueError("Old .xls files aren't supported. Open it in Excel and use Save As -> Excel Workbook (.xlsx).")
        if ext not in (".xlsx", ".xlsm"):
            raise ValueError(f"Unsupported file type '{ext}'. Use an Excel .xlsx/.xlsm file or a Google Sheets link.")
        if not self.path.exists():
            raise FileNotFoundError(str(self.path))
        self._openpyxl = openpyxl
        self._values = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
        self.wb = None
        self.display_name = self.path.name
        self.backup_path: Path | None = None

    # ------------------------------------------------------------ reading
    def tab_names(self) -> list[str]:
        return [ws.title for ws in self._values.worksheets]

    def read_tab(self, name: str) -> list[list]:
        ws = self._values[name]
        return [list(r) for r in ws.iter_rows(values_only=True)]

    # ------------------------------------------------------------ writing
    def begin_write(self) -> list[str]:
        """Load the editable copy; return warnings about content openpyxl cannot round-trip."""
        self._values.close()
        warnings = _unsupported_features(self.path)
        self.wb = self._openpyxl.load_workbook(self.path, keep_vba=self.path.suffix.lower() == ".xlsm")
        return warnings

    def pod_folder(self) -> Path:
        return self.path.parent / f"{self.path.stem} - POD"

    def pod_link(self, path: Path) -> str:
        return file_link(path, self.path.parent)

    def write_status_tab(self, name: str, headers: list[str], rows: list[list[OutCell]], widths: dict[str, int],
                         header_notes: dict[str, str] | None = None) -> None:
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.hyperlink import Hyperlink

        wb = self.wb
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name)
        head_fill = PatternFill("solid", fgColor="1F4E78")
        for c, h in enumerate(headers, 1):
            cell = ws.cell(1, c, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if header_notes and h in header_notes:
                cell.comment = Comment(header_notes[h][:2000], "Tracking Check", width=420, height=220)
        ws.row_dimensions[1].height = 32

        for r, row in enumerate(rows, 2):
            for c, oc in enumerate(row, 1):
                cell = ws.cell(r, c, _xl(oc.value))
                if oc.url:
                    cell.hyperlink = oc.url
                    cell.style = "Hyperlink"
                elif oc.goto:
                    tab, gr, gc = oc.goto
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"{quote_tab(tab)}!{_a1(gr, gc)}",
                                               tooltip="Go to this tracking number on the original tab")
                    cell.style = "Hyperlink"
                if isinstance(oc.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                elif isinstance(oc.value, date):
                    cell.number_format = "yyyy-mm-dd"
                if oc.fill:
                    cell.fill = PatternFill("solid", fgColor=oc.fill)
        for c, h in enumerate(headers, 1):
            ws.column_dimensions[_col(c)].width = widths.get(h, 16)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{_col(len(headers))}{max(1, len(rows) + 1)}"

    def link_source_cells(self, src_tab: str, links: list[tuple[int, int, int]], status_tab: str) -> int:
        from openpyxl.cell.cell import MergedCell
        from openpyxl.worksheet.hyperlink import Hyperlink

        ws = self.wb[src_tab]
        done = 0
        for row, col, target_row in links:
            cell = ws.cell(row, col)
            if isinstance(cell, MergedCell):
                continue
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"{quote_tab(status_tab)}!A{target_row}",
                                       tooltip="Show tracking status")
            f = copy(cell.font)
            f.color = "0563C1"
            f.underline = "single"
            cell.font = f
            done += 1
        return done

    def save(self) -> str:
        bdir = workspace_dir() / "backups"
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        self.backup_path = bdir / f"{self.path.stem}_{stamp}{self.path.suffix}"
        shutil.copy2(self.path, self.backup_path)
        try:
            self.wb.save(self.path)
        except PermissionError:
            raise PermissionError(f"Can't save {self.path.name} - it is probably open in Excel. Close it and run again.")
        return str(self.path)


def _unsupported_features(path: Path) -> list[str]:
    checks = {
        "xl/slicers/": "slicers", "xl/timelines/": "timelines", "xl/ctrlProps/": "form controls",
        "xl/activeX/": "ActiveX controls", "xl/model/": "a Power Pivot data model",
    }
    found = set()
    try:
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                for prefix, label in checks.items():
                    if n.startswith(prefix):
                        found.add(label)
                if n.startswith("xl/drawings/drawing") and n.endswith(".xml"):
                    if b"<xdr:sp" in z.read(n):
                        found.add("drawn shapes/text boxes")
    except zipfile.BadZipFile:
        pass
    return [f"This workbook contains {x}, which the Excel library may drop when saving. "
            f"A backup copy is kept." for x in sorted(found)]


_ILLEGAL = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


def _xl(v):
    if isinstance(v, str):
        return _ILLEGAL.sub("", v)
    return v


def _col(n: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def _a1(row: int, col: int) -> str:
    return f"{_col(col)}{row}"
