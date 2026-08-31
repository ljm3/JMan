from __future__ import annotations

from pathlib import Path

from .base import ExtractResult


def extract_docx(path: Path, cfg) -> ExtractResult:
    from docx import Document  # python-docx
    doc = Document(str(path))
    paras = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    tbl_text = []
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                tbl_text.append(" | ".join(cells))
    cp = doc.core_properties
    meta = {
        "paragraphs": len(paras),
        "tables": len(doc.tables),
        "title": cp.title or "",
        "author": cp.author or "",
        "created": cp.created.isoformat() if cp.created else "",
        "modified": cp.modified.isoformat() if cp.modified else "",
        "last_modified_by": cp.last_modified_by or "",
    }
    body = "\n".join(paras + tbl_text)
    return ExtractResult(text=body, kind="document", extractor="python-docx", meta=meta,
                         tools=[("docx", "read .docx documents")])


def _cellval(v):
    """JSON-friendly scalar for a spreadsheet cell."""
    import datetime as _dt
    if v is None:
        return ""
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, _dt.timedelta):
        return v.total_seconds() / 60.0            # minutes - handy for time studies
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return v.isoformat()
    return str(v)


def _norm_header(s: str) -> str:
    return " ".join(str(s).split())


def _scan_sheet_table(name: str, rows: list[tuple]) -> dict | None:
    """Heuristically find the header row and return a structured table + the
    label:value 'context' cells around it (DATE:, RESOURCE:, Total Productive
    Time:, ...).  Generic - works for any reasonably tabular sheet."""
    if not rows:
        return None
    n = len(rows)

    def score_header(r):
        cells = [c for c in r if c not in (None, "")]
        strs = [c for c in cells if isinstance(c, str) and 1 <= len(c.strip()) <= 40]
        return len(strs) if len(strs) >= 3 and len(strs) >= 0.6 * len(cells) else 0

    best_i, best = -1, 0
    for i in range(min(n, 40)):
        s = score_header(rows[i])
        # header should have data rows under it
        if s > best and i + 1 < n and any(c not in (None, "") for c in rows[i + 1]):
            best, best_i = s, i
    if best_i < 0:
        return None

    header = [_norm_header(c) if c not in (None, "") else f"col{j+1}"
              for j, c in enumerate(rows[best_i])]
    # trim trailing empty-ish columns
    while header and header[-1].startswith("col"):
        header.pop()
    if len(header) < 2:
        return None
    ncol = len(header)

    data: list[dict] = []
    for r in rows[best_i + 1:]:
        vals = list(r)[:ncol]
        if not any(v not in (None, "") for v in vals):
            continue
        rec = {header[j]: _cellval(vals[j]) if j < len(vals) else "" for j in range(ncol)}
        data.append(rec)
        if len(data) >= 8000:
            break

    context: dict[str, object] = {}
    scan_upto = min(n, best_i + 6) if data else n
    for i in range(n):
        if best_i < i < scan_upto and i > best_i:
            pass
        r = rows[i]
        for j, c in enumerate(r):
            if isinstance(c, str) and c.strip().endswith(":") and len(c.strip()) <= 48:
                label = c.strip().rstrip(":").strip()
                val = ""
                for k in range(j + 1, min(len(r), j + 4)):
                    if r[k] not in (None, ""):
                        val = _cellval(r[k])
                        break
                if label and label not in context:
                    context[label] = val
    return {"sheet": name, "header_row": best_i + 1, "columns": header,
            "row_count": len(data), "rows": data, "context": context}


def extract_xlsx(path: Path, cfg) -> ExtractResult:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    sheet_info = []
    total_cells = 0
    tables: list[dict] = []
    merged_context: dict[str, object] = {}
    budget = 12000
    for ws in wb.worksheets:
        raw_rows: list[tuple] = []
        rows = 0
        for row in ws.iter_rows(values_only=True):
            vals = [("" if v is None else str(v)) for v in row]
            if any(vals):
                lines.append(" | ".join(vals))
                total_cells += sum(1 for v in vals if v)
            if rows < 8200:
                raw_rows.append(row)
            rows += 1
            if rows > 20000:
                break
        sheet_info.append({"name": ws.title, "rows": rows,
                           "dims": ws.calculate_dimension() if hasattr(ws, "calculate_dimension") else ""})
        if budget > 0:
            tbl = _scan_sheet_table(ws.title, raw_rows)
            if tbl:
                take = min(len(tbl["rows"]), budget)
                tbl["rows"] = tbl["rows"][:take]
                tbl["row_count"] = take
                budget -= take
                tables.append(tbl)
                for k, v in tbl["context"].items():
                    merged_context.setdefault(k, v)
    wb.close()
    props = wb.properties
    meta = {
        "sheets": [s["name"] for s in sheet_info],
        "sheet_detail": sheet_info,
        "non_empty_cells": total_cells,
        "creator": getattr(props, "creator", "") or "",
        "created": props.created.isoformat() if getattr(props, "created", None) else "",
        "modified": props.modified.isoformat() if getattr(props, "modified", None) else "",
    }
    if tables:
        meta["record_tables"] = tables
        meta["cell_context"] = merged_context
        primary = max(tables, key=lambda t: t["row_count"])
        meta["rows"] = primary["row_count"]
        meta["columns"] = len(primary["columns"])
    return ExtractResult(text="\n".join(lines), kind="spreadsheet", extractor="openpyxl",
                         meta=meta, tools=[("openpyxl", "read .xlsx spreadsheets")])


def extract_pptx(path: Path, cfg) -> ExtractResult:
    from pptx import Presentation
    prs = Presentation(str(path))
    chunks: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts = [f"--- slide {i} ---"]
        for shp in slide.shapes:
            if shp.has_text_frame:
                for p in shp.text_frame.paragraphs:
                    line = "".join(r.text for r in p.runs) or p.text
                    if line.strip():
                        parts.append(line)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            note = slide.notes_slide.notes_text_frame.text
            if note.strip():
                parts.append(f"[notes] {note}")
        chunks.append("\n".join(parts))
    cp = prs.core_properties
    meta = {
        "slides": len(prs.slides),
        "title": cp.title or "",
        "author": cp.author or "",
        "created": cp.created.isoformat() if cp.created else "",
        "modified": cp.modified.isoformat() if cp.modified else "",
    }
    return ExtractResult(text="\n\n".join(chunks), kind="presentation", extractor="python-pptx",
                         meta=meta, tools=[("pptx", "read .pptx decks")])


def extract_legacy(path: Path, cfg) -> ExtractResult:
    return ExtractResult(
        kind="document", extractor="legacy-office (unsupported)", ok=False,
        error=f"legacy binary Office format '{path.suffix}' is not parsed; "
              "re-save as the modern XML format (.docx/.xlsx/.pptx) to include it",
        meta={"skipped": True},
    )
