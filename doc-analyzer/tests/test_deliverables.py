"""The choose-format result documents: section builders + the three renderers."""
from __future__ import annotations

from doc_analyzer.report.formats import (
    FMT_LABEL,
    LABEL_TO_FMT,
    Section,
    build_activity_sections,
    build_synopsis_sections,
    build_tools_sections,
    render_presentation,
    render_spreadsheet,
    render_word,
    resolve_fmt,
)


class _Cfg:
    default_output_format = "word"


def _corpus():
    return {
        "generated_utc": "2026-08-30T00:00:00+00:00",
        "source": "Local folder: /tmp/docs",
        "objective": "Find the money",
        "root": "/tmp/docs",
        "doc_count": 2,
        "files_found": 3,
        "files_skipped": 1,
        "total_words": 120,
        "total_bytes": 4096,
        "type_distribution": [("prose", 2)],
        "ext_distribution": [(".txt", 2)],
        "keywords": ["invoice", "payment", "acme"],
        "timeline": [{"date": "2026-03-01", "doc": "a.txt", "kind": "prose"}],
    }


def _records():
    return [
        {"rel": "a.txt", "kind": "prose", "extractor": "text", "size": 2048, "words": 80,
         "ok": True, "error": "", "best_date": "2026-03-01",
         "keywords": ["invoice", "acme"], "summary": ["Acme invoice for March."],
         "entities": {"money": ["$1,091.88"], "organizations": ["Acme"]}},
        {"rel": "b.txt", "kind": "prose", "extractor": "text", "size": 2048, "words": 40,
         "ok": True, "error": "", "best_date": "2026-03-04", "keywords": ["payment"],
         "summary": ["Payment received."], "entities": {}},
        {"rel": "c.bin", "kind": "binary", "extractor": "skipped", "size": 10, "words": 0,
         "ok": False, "error": "no extractable text", "best_date": "", "keywords": [],
         "summary": [], "entities": {}},
    ]


def _clusters():
    return [{"id": 1, "size": 2, "label": "invoice", "shared_terms": ["invoice", "acme"],
             "kinds": ["prose"], "representative": "a.txt", "docs": ["a.txt", "b.txt"]},
            {"id": 2, "size": 1, "label": "", "shared_terms": [], "kinds": ["prose"],
             "representative": "d.txt", "docs": ["d.txt"]}]


def _phrases():
    return [{"phrase": "receive payment", "category": "Finance & billing", "verb": "receive",
             "occurrences": 3, "doc_count": 2, "docs": ["a.txt", "b.txt"]}]


def _categories():
    return [{"category": "Finance & billing", "doc_count": 2, "mentions": 5,
             "verbs": ["invoice", "pay"], "docs": ["a.txt", "b.txt"]}]


def _ledger_dict():
    return {
        "session_started_utc": "2026-08-30T00:00:00+00:00",
        "generated_utc": "2026-08-30T00:01:00+00:00",
        "answers": {"What do you want to ascertain from this analysis?": "Find the money"},
        "generated_files": [{"deliverable": "synopsis", "format": "word", "file": "synopsis.docx"}],
        "tool_count": 2,
        "tools": [
            {"category": "runtime", "name": "python", "version": "3.12.0",
             "detail": "/usr/bin/python", "purposes": ["interpreter"]},
            {"category": "python-package", "name": "openpyxl", "version": "3.1.2",
             "detail": None, "purposes": ["render the spreadsheet deliverable"]},
        ],
    }


def test_label_maps_round_trip():
    for fmt, label in FMT_LABEL.items():
        assert LABEL_TO_FMT[label] == fmt


def test_resolve_fmt_prefers_run_opts_then_cfg_then_default():
    class RO:
        formats = {"synopsis": "presentation"}

    assert resolve_fmt(RO(), _Cfg(), "synopsis") == "presentation"
    assert resolve_fmt(RO(), _Cfg(), "activity") == "word"          # cfg default
    assert resolve_fmt(None, _Cfg(), "tools") == "word"
    _Cfg.default_output_format = "bogus"
    assert resolve_fmt(None, _Cfg(), "tools") == "word"             # invalid -> word
    _Cfg.default_output_format = "word"


def test_synopsis_sections_include_objective_and_tables():
    secs = build_synopsis_sections(_corpus(), _records(), _records()[:2], None,
                                   "Find the money", detailed=True)
    assert secs[0].heading == "Objective"
    assert "Find the money" in secs[0].paragraphs[0]
    headings = [s.heading for s in secs]
    assert "Overview" in headings and "Per-document overview" in headings
    assert "Skipped / errored files" in headings


def test_activity_and_tools_sections():
    act = build_activity_sections(_corpus(), _clusters(), _phrases(), _categories(), "")
    assert act[0].paragraphs[0] == "(not specified)"
    assert any(s.table for s in act)

    tools = build_tools_sections(_ledger_dict(), "Find the money")
    joined = [s.heading for s in tools]
    assert "Analysis answers" in joined and "Generated files" in joined
    assert "Tools & versions" in joined


def _sections():
    return [
        Section("Intro", 1, paragraphs=["Line one.", "Line two."], bullets=["a", "b"]),
        Section("Numbers", 1, table=(["Name", "Count"], [["x", 1], ["y", 2]])),
    ]


def test_render_word_opens(tmp_path):
    import docx

    p = render_word("Title", "sub", _sections(), tmp_path / "d.docx")
    doc = docx.Document(str(p))
    assert len(doc.tables) == 1
    assert any("Intro" in para.text for para in doc.paragraphs)


def test_render_spreadsheet_opens(tmp_path):
    import openpyxl

    p = render_spreadsheet("Title", "sub", _sections(), tmp_path / "d.xlsx")
    wb = openpyxl.load_workbook(str(p))
    assert "Overview" in wb.sheetnames
    assert any(name.startswith("Numbers") for name in wb.sheetnames)


def test_render_presentation_opens(tmp_path):
    from pptx import Presentation

    p = render_presentation("Title", "sub", _sections(), tmp_path / "d.pptx")
    prs = Presentation(str(p))
    assert len(prs.slides) >= 3  # title + text slide + table slide


def test_renderers_survive_long_table(tmp_path):
    big = [Section("Big", 1, table=(["i", "v"], [[i, f"row {i}"] for i in range(200)]))]
    render_word("T", "", big, tmp_path / "b.docx")
    render_spreadsheet("T", "", big, tmp_path / "b.xlsx")
    render_presentation("T", "", big, tmp_path / "b.pptx")
    for name in ("b.docx", "b.xlsx", "b.pptx"):
        assert (tmp_path / name).stat().st_size > 0
